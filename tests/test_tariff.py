"""Tests for azoic.tariff: multiplicative-tariff export for a fitted RiskGLM.

M6 acceptance: the tariff's recalibrated base reproduces the portfolio total
observed claim amount exactly, and the structural (non-recalibrated) tariff's
``apply_tariff`` reproduces ``RiskGLM.predict`` up to float round-off.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from azoic.metrics import mean_tweedie_deviance
from azoic.models import RiskGBM, RiskGLM
from azoic.tariff import (
    apply_tariff,
    distill_gbm,
    export_tariff,
    extract_tariff,
    recalibrate_for_total,
)
from tests.conftest import make_synthetic_portfolio


def _fit_glm(n: int = 2000, seed: int = 42, **glm_kwargs) -> tuple[RiskGLM, pd.DataFrame]:
    df = make_synthetic_portfolio(n=n, seed=seed)
    features = ["driver_age", "vehicle_age", "region", "vehicle_brand"]
    X = df[features + ["exposure"]]
    y = (df["claim_amount"] / df["exposure"]).to_numpy(dtype=float)
    kwargs = {
        "family": "tweedie",
        "link": "log",
        "exposure_col": "exposure",
        "tweedie_power": 1.5,
        **glm_kwargs,
    }
    glm = RiskGLM(**kwargs)
    glm.fit(X, y)
    return glm, df


def _fit_gbm(n: int = 3000) -> tuple[RiskGBM, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    df = make_synthetic_portfolio(n=n, seed=42)
    X = df[["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]]
    X_fit, X_validation = X.iloc[: 4 * n // 5], X.iloc[4 * n // 5 :]
    y_fit = (df["claim_amount"].iloc[: len(X_fit)] / df["exposure"].iloc[: len(X_fit)]).to_numpy(
        dtype=float
    )
    teacher = RiskGBM(
        objective="tweedie",
        exposure_col="exposure",
        tweedie_variance_power=1.5,
        n_estimators=80,
        num_leaves=15,
        learning_rate=0.05,
        random_state=42,
    ).fit(X_fit, y_fit)
    return teacher, df, X_fit, X_validation


def test_distill_gbm_beats_constant_and_preserves_teacher_total() -> None:
    teacher, _, X_fit, X_validation = _fit_gbm()

    student = distill_gbm(teacher, X_fit, X_validation)
    teacher_fit = teacher.predict(X_fit)
    teacher_validation = teacher.predict(X_validation)
    student_validation = student.predict(X_validation)
    exposure_fit = X_fit["exposure"].to_numpy(dtype=float)
    exposure_validation = X_validation["exposure"].to_numpy(dtype=float)
    constant = np.full(
        len(X_validation),
        np.average(teacher_fit, weights=exposure_fit),
    )
    student_deviance = mean_tweedie_deviance(
        teacher_validation,
        student_validation,
        sample_weight=exposure_validation,
        power=1.5,
    )
    constant_deviance = mean_tweedie_deviance(
        teacher_validation,
        constant,
        sample_weight=exposure_validation,
        power=1.5,
    )

    assert student.family == "tweedie"
    assert student.link == "log"
    assert student.tweedie_power == 1.5
    assert student_deviance < constant_deviance
    assert student.distillation_metrics_["teacher_student_deviance"] == pytest.approx(
        student_deviance
    )
    assert student.distillation_metrics_["student_teacher_total_ratio"] == pytest.approx(
        1.0, rel=0.05
    )
    assert np.allclose(
        apply_tariff(extract_tariff(student), X_validation),
        student_validation,
    )


@pytest.mark.parametrize("objective", ["tweedie", "poisson", "gamma"])
def test_distill_gbm_maps_positive_objectives(objective: str) -> None:
    rng = np.random.default_rng(42)
    X = pd.DataFrame(
        {
            "x": rng.normal(size=500),
            "exposure": rng.uniform(0.1, 1.0, size=500),
        }
    )
    y = np.exp(1.0 + 0.2 * X["x"].to_numpy())
    teacher = RiskGBM(
        objective=objective,
        exposure_col="exposure",
        tweedie_variance_power=1.3,
        n_estimators=20,
        random_state=42,
    ).fit(X.iloc[:400], y[:400])

    student = distill_gbm(teacher, X.iloc[:400], X.iloc[400:])

    assert student.family == objective
    assert student.link == "log"
    assert student.tweedie_power == (1.3 if objective == "tweedie" else None)


def test_distill_gbm_rejects_invalid_teacher_inputs(monkeypatch) -> None:
    _, _, X_fit, X_validation = _fit_gbm(n=1000)
    with pytest.raises(ValueError, match="fitted"):
        distill_gbm(RiskGBM(exposure_col="exposure"), X_fit, X_validation)

    y_fit = np.ones(len(X_fit))
    unsupported = RiskGBM(
        objective="regression",
        exposure_col="exposure",
        n_estimators=5,
    ).fit(X_fit, y_fit)
    with pytest.raises(ValueError, match="supports only"):
        distill_gbm(unsupported, X_fit, X_validation)

    teacher, _, X_fit, X_validation = _fit_gbm(n=1000)
    for invalid in (0.0, np.nan):
        invalid_exposure = X_validation.copy()
        invalid_exposure.iloc[0, invalid_exposure.columns.get_loc("exposure")] = invalid
        with pytest.raises(ValueError, match="positive and finite"):
            distill_gbm(teacher, X_fit, invalid_exposure)
    with pytest.raises(ValueError, match="separate"):
        distill_gbm(teacher, X_fit, X_fit)

    for invalid in (0.0, np.inf):
        monkeypatch.setattr(teacher, "predict", lambda X, value=invalid: np.full(len(X), value))
        with pytest.raises(ValueError, match="predictions must be positive and finite"):
            distill_gbm(teacher, X_fit, X_validation)


# ---------------------------------------------------------------------------
# extract_tariff
# ---------------------------------------------------------------------------


def test_extract_tariff_unfitted_raises() -> None:
    with pytest.raises(ValueError, match="fitted"):
        extract_tariff(RiskGLM(family="tweedie", link="log"))


def test_extract_tariff_non_log_link_raises(tmp_path: Path) -> None:
    glm, _ = _fit_glm(link="identity", family="normal")
    with pytest.raises(ValueError, match="log-link"):
        extract_tariff(glm)


def test_extract_tariff_structure_numeric_and_categorical() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    assert set(t.keys()) == {"base_rate", "reference", "numeric", "categorical", "mapping"}
    # Numeric features carry their raw coefficient.
    assert "driver_age" in t["numeric"] and "vehicle_age" in t["numeric"]
    # Categorical features carry one entry per observed level.
    assert set(t["categorical"]["region"].keys()) == {"rural", "suburban", "urban"}
    assert set(t["categorical"]["vehicle_brand"].keys()) == {"A", "B", "C", "D"}
    # Every reference level has factor exactly 1.0.
    for feat, ref in t["reference"].items():
        assert t["categorical"][feat][ref] == pytest.approx(1.0)


def test_extract_tariff_default_reference_is_first_fitted_level() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    assert t["reference"]["region"] == "rural"
    assert t["reference"]["vehicle_brand"] == "A"


def test_extract_tariff_reference_override_respected() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm, reference={"region": "urban", "vehicle_brand": "D"})
    assert t["reference"]["region"] == "urban"
    assert t["reference"]["vehicle_brand"] == "D"
    assert t["categorical"]["region"]["urban"] == pytest.approx(1.0)
    assert t["categorical"]["vehicle_brand"]["D"] == pytest.approx(1.0)


def test_extract_tariff_reference_unknown_level_raises() -> None:
    glm, _ = _fit_glm()
    with pytest.raises(ValueError, match="not in feature"):
        extract_tariff(glm, reference={"region": "atlantis"})


def test_extract_tariff_mapping_frame_lists_features_and_roles() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    m = t["mapping"]
    assert list(m.columns) == ["feature", "role", "dtype", "levels", "n_levels", "reference_level"]
    # Each input feature appears exactly once (deduplicated across levels).
    assert sorted(m["feature"]) == ["driver_age", "region", "vehicle_age", "vehicle_brand"]
    roles = dict(zip(m["feature"], m["role"], strict=True))
    assert roles["driver_age"] == "numeric"
    assert roles["region"] == "categorical"
    # Categorical rows expose their levels + the chosen reference.
    region_row = m[m["feature"] == "region"].iloc[0]
    assert "rural" in str(region_row["levels"]) and "urban" in str(region_row["levels"])
    assert region_row["reference_level"] == "rural"
    assert int(region_row["n_levels"]) == 3


# ---------------------------------------------------------------------------
# apply_tariff roundtrip to glm.predict (no recalibration)
# ---------------------------------------------------------------------------


def test_apply_tariff_roundtrips_to_glm_predict() -> None:
    glm, df = _fit_glm()
    t = extract_tariff(glm)
    feats = ["driver_age", "vehicle_age", "region", "vehicle_brand"]
    applied = apply_tariff(t, df[feats])
    predicted = np.asarray(glm.predict(df[feats]), dtype=float)
    assert np.allclose(applied, predicted, rtol=1e-6, atol=1e-3)


@pytest.mark.parametrize(
    ("numeric", "values", "expected"),
    [
        ({"x": 1000.0}, {"x": [0.001]}, np.e),
        ({"x": 1000.0, "z": -1000.0}, {"x": [1.0], "z": [1.0]}, 1.0),
    ],
)
def test_apply_tariff_sums_before_exponentiating(numeric, values, expected) -> None:
    tariff = {"base_rate": 1.0, "numeric": numeric, "categorical": {}}
    assert apply_tariff(tariff, pd.DataFrame(values)) == pytest.approx([expected])


def test_extract_tariff_folds_reference_before_exponentiating() -> None:
    X = pd.DataFrame({"cat": pd.Categorical(["a", "b"] * 10)})
    glm = RiskGLM(family="poisson", link="log").fit(X, np.tile([1.0, 2.0], 10))
    glm.backend_.intercept_ = 1000.0
    glm.backend_.coef_[:] = [-1000.0, -999.0]
    tariff = extract_tariff(glm)
    assert tariff["base_rate"] == pytest.approx(1.0)
    np.testing.assert_allclose(apply_tariff(tariff, X), glm.predict(X))


@pytest.mark.parametrize("bad", [np.nan, np.inf, -np.inf])
def test_extract_tariff_rejects_nonfinite_coefficients(bad) -> None:
    glm, _ = _fit_glm(n=100)
    glm.backend_.coef_[0] = bad
    with pytest.raises(ValueError, match="finite"):
        extract_tariff(glm)


@pytest.mark.parametrize("coefficient", [-1000.0, 1000.0])
@pytest.mark.parametrize("part", ["base", "categorical", "workbook"])
def test_tariff_rejects_unrepresentable_factors(tmp_path, coefficient, part) -> None:
    X = pd.DataFrame({"x": [0.001] * 20, "cat": pd.Categorical(["a", "b"] * 10)})
    glm = RiskGLM(family="poisson", link="log").fit(X, np.tile([1.0, 2.0], 10))
    glm.backend_.intercept_ = coefficient if part == "base" else 0.0
    glm.backend_.coef_[:] = 0.0
    glm.backend_.coef_[0 if part == "workbook" else 2] = coefficient if part != "base" else 0.0
    out = tmp_path / "invalid.xlsx"
    if part == "workbook":
        tariff = extract_tariff(glm)
        np.testing.assert_allclose(apply_tariff(tariff, X), glm.predict(X))
    with pytest.raises(ValueError, match="positive and finite"):
        if part == "workbook":
            export_tariff(glm, out, recalibrate=False)
        else:
            extract_tariff(glm)
    assert not out.exists()


@pytest.mark.parametrize(
    "part, bad",
    [(part, bad) for part in ["base", "numeric", "categorical"] for bad in [np.nan, np.inf]]
    + [("base", -1.0), ("categorical", -1.0), ("categorical", 0.0)],
)
def test_apply_tariff_rejects_invalid_parameters(bad, part) -> None:
    tariff = {"base_rate": 1.0, "numeric": {"x": 0.0}, "categorical": {"cat": {"a": 1.0}}}
    if part == "base":
        tariff["base_rate"] = bad
    elif part == "numeric":
        tariff["numeric"]["x"] = bad
    else:
        tariff["categorical"]["cat"]["a"] = bad
    with pytest.raises(ValueError, match="finite"):
        apply_tariff(tariff, pd.DataFrame({"x": [0.001], "cat": ["a"]}))


@pytest.mark.parametrize("coefficient", [-1000.0, 1000.0])
def test_apply_tariff_rejects_unrepresentable_rates(coefficient) -> None:
    tariff = {"base_rate": 1.0, "numeric": {"x": coefficient}, "categorical": {}}
    with pytest.raises(ValueError, match="positive and finite"):
        apply_tariff(tariff, pd.DataFrame({"x": [1.0]}))


def test_zero_total_recalibration_keeps_zero_rates(tmp_path) -> None:
    glm, df = _fit_glm(n=100)
    X = df[list(glm.feature_names_in_)]
    out = export_tariff(
        glm, tmp_path / "zero.xlsx", X=X, y=np.zeros(len(X)), exposure_col="exposure"
    )
    tariff = extract_tariff(glm)
    tariff["base_rate"] = pd.read_excel(out, sheet_name="base_rate").loc[0, "base_rate"]
    np.testing.assert_array_equal(apply_tariff(tariff, X), np.zeros(len(X)))
    with pytest.raises(ValueError, match="unknown levels"):
        apply_tariff(tariff, X.assign(region="unknown"))


@pytest.mark.parametrize(
    "base, predicted, observed, expected",
    [
        (1e300, 1e300, 1e-300, 1e-300),
        (1e-300, 1e-300, 1e300, 1e300),
    ],
)
def test_recalibration_avoids_intermediate_overflow(base, predicted, observed, expected) -> None:
    result = recalibrate_for_total(
        {"base_rate": base}, predicted_total=predicted, observed_total=observed
    )
    assert result / expected == pytest.approx(1.0)


@pytest.mark.parametrize("base, predicted, observed", [(1e300, 1e-300, 1.0), (1e-300, 1e300, 1.0)])
def test_recalibration_rejects_unrepresentable_base(base, predicted, observed) -> None:
    with pytest.raises(ValueError, match="positive and finite"):
        recalibrate_for_total(
            {"base_rate": base}, predicted_total=predicted, observed_total=observed
        )


@pytest.mark.parametrize("pipeline", [False, True])
def test_export_tariff_preserves_literal_text_and_typed_labels(tmp_path, pipeline) -> None:
    from openpyxl import load_workbook
    from sklearn.pipeline import Pipeline

    from azoic.preprocessing import AutoGrouper

    X = pd.DataFrame(
        {
            "=feature": pd.Categorical(
                ["=1+1", "#N/A"] * 20, categories=["=1+1", "#N/A"], ordered=True
            ),
            "number": pd.Categorical([2, 1] * 20),
        }
    )
    model = RiskGLM(family="poisson", link="log")
    if pipeline:
        model = Pipeline(
            [("grouper", AutoGrouper(cols=["=feature"], strategy="rare")), ("model", model)]
        )
    else:
        X["flag"] = pd.Categorical([True, False] * 20)
    model.fit(X, np.tile([1.0, 2.0], 20))
    out = export_tariff(
        model, tmp_path / "literal.xlsx", reference={"=feature": "=1+1"}, recalibrate=False
    )
    workbook = load_workbook(out)
    assert workbook.sheetnames == ["base_rate", "factors", "mappings"]
    rows = list(workbook["factors"].iter_rows(min_row=2))
    for row in rows:
        feature, level = row[:2]
        if feature.value == "=feature":
            assert feature.data_type == "s"
            assert level.value in {"=1+1", "#N/A", "Other"}
            assert level.data_type == "s"
        elif feature.value == "flag":
            assert type(level.value) is bool and level.data_type == "b"
        else:
            assert type(level.value) is int and level.data_type == "n"
    assert any(cell.value == "=1+1" for row in rows for cell in row)
    mapping_rows = list(workbook["mappings"].iter_rows(min_row=2))
    assert any(row[1].value == "grouped" for row in mapping_rows) == pipeline
    assert any(row[5].value == "=1+1" for row in mapping_rows)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.data_type == "s"


def test_extract_tariff_preserves_typed_ordered_levels() -> None:
    rng = np.random.default_rng(42)
    intervals = [pd.Interval(0, 1), pd.Interval(1, 2)]
    X = pd.DataFrame(
        {
            "ordered_int": pd.Categorical(
                rng.choice([2, 1, 3], 300), categories=[2, 1, 3], ordered=True
            ),
            "flag": pd.Categorical(
                rng.choice([True, False], 300), categories=[True, False], ordered=True
            ),
            "band": pd.Categorical(rng.choice(intervals, 300), categories=intervals, ordered=True),
            "x": rng.normal(size=300),
        }
    )
    y = np.exp(0.1 + 0.2 * X["x"].to_numpy() + rng.normal(scale=0.1, size=len(X)))
    glm = RiskGLM(family="gamma", link="log", alpha=0.1).fit(X, y)

    tariff = extract_tariff(glm, reference={"ordered_int": 1})

    assert list(tariff["categorical"]["ordered_int"]) == [2, 1, 3]
    assert list(tariff["categorical"]["flag"]) == [True, False]
    assert list(tariff["categorical"]["band"]) == intervals
    assert {type(level) for level in tariff["categorical"]["ordered_int"]} == {int}
    assert {type(level) for level in tariff["categorical"]["flag"]} == {bool}
    assert {type(level) for level in tariff["categorical"]["band"]} == {pd.Interval}
    assert tariff["reference"] == {"ordered_int": 1, "flag": True, "band": intervals[0]}
    assert tariff["categorical"]["ordered_int"][1] == pytest.approx(1.0)
    assert np.allclose(apply_tariff(tariff, X), glm.predict(X))
    with pytest.raises(ValueError, match="not in feature"):
        extract_tariff(glm, reference={"flag": 1})


def test_apply_tariff_requires_dataframe() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    with pytest.raises(TypeError, match="DataFrame"):
        apply_tariff(t, np.zeros((3, 4)))


def test_apply_tariff_missing_numeric_column_raises() -> None:
    glm, df = _fit_glm()
    t = extract_tariff(glm)
    X = df.drop(columns=["driver_age"])
    with pytest.raises(KeyError, match="driver_age"):
        apply_tariff(t, X[["vehicle_age", "region", "vehicle_brand"]])


def test_apply_tariff_missing_categorical_column_raises() -> None:
    glm, df = _fit_glm()
    t = extract_tariff(glm)
    X = df.drop(columns=["region"])
    with pytest.raises(KeyError, match="region"):
        apply_tariff(t, X[["driver_age", "vehicle_age", "vehicle_brand"]])


def test_apply_tariff_unknown_categorical_level_raises() -> None:
    glm, df = _fit_glm()
    tariff = extract_tariff(glm)
    X = df.head(3).copy()
    X.loc[:, "region"] = "atlantis"
    with pytest.raises(ValueError, match="unknown levels"):
        apply_tariff(tariff, X)


# ---------------------------------------------------------------------------
# recalibrate_for_total
# ---------------------------------------------------------------------------


def test_recalibrate_for_total_reproduces_observed_total() -> None:
    glm, df = _fit_glm()
    t = extract_tariff(glm)
    feats = ["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]
    exp = df["exposure"].to_numpy(float)
    cam = df["claim_amount"].to_numpy(float)
    pred = np.asarray(glm.predict(df[feats]), dtype=float)
    pred_total = float((pred * exp).sum())
    obs_total = float(cam.sum())

    new_base = recalibrate_for_total(t, predicted_total=pred_total, observed_total=obs_total)
    t2 = dict(t)
    t2["base_rate"] = new_base
    recal = apply_tariff(t2, df[["driver_age", "vehicle_age", "region", "vehicle_brand"]])
    recal_total = float((recal * exp).sum())
    assert recal_total == pytest.approx(obs_total, rel=1e-7)


def test_recalibrate_for_total_nonpositive_predicted_raises() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    with pytest.raises(ValueError, match="positive"):
        recalibrate_for_total(t, predicted_total=0.0, observed_total=100.0)


def test_recalibrate_for_total_negative_observed_raises() -> None:
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    with pytest.raises(ValueError, match="non-negative"):
        recalibrate_for_total(t, predicted_total=100.0, observed_total=-1.0)


@pytest.mark.parametrize(
    ("pred", "obs"), [(np.nan, 100.0), (np.inf, 100.0), (100.0, np.nan), (100.0, np.inf)]
)
def test_recalibrate_for_total_nonfinite_totals_raise(pred, obs) -> None:
    """NaN/inf slip past the <= 0 / < 0 comparisons and silently wrote a NaN
    (or zero, for inf exposure) base rate into the workbook (M33 F2)."""
    glm, _ = _fit_glm()
    t = extract_tariff(glm)
    with pytest.raises(ValueError, match="finite"):
        recalibrate_for_total(t, predicted_total=pred, observed_total=obs)


@pytest.mark.parametrize("bad", [np.nan, np.inf, -1.0, 0.001])
def test_export_tariff_recalibration_rejects_invalid_exposure(tmp_path, bad) -> None:
    glm, df = _fit_glm()
    feats = ["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]
    X = df[feats].copy()
    X.loc[df.index[0], "exposure"] = bad
    match = "1/366" if bad == 0.001 else "finite|non-negative"
    with pytest.raises(ValueError, match=match):
        export_tariff(glm, tmp_path / "t.xlsx", X=X, y=df["claim_amount"], exposure_col="exposure")


@pytest.mark.parametrize("bad", [np.nan, np.inf, -1.0])
def test_export_tariff_recalibration_rejects_invalid_y(tmp_path, bad) -> None:
    glm, df = _fit_glm()
    feats = ["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]
    y = df["claim_amount"].astype(float).copy()
    y.iloc[0] = bad
    with pytest.raises(ValueError, match="NaN|infinity|finite|non-negative"):
        export_tariff(glm, tmp_path / "t.xlsx", X=df[feats], y=y, exposure_col="exposure")


def test_export_tariff_recalibration_rejects_mismatched_y_length(tmp_path) -> None:
    glm, df = _fit_glm()
    feats = ["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]
    with pytest.raises(ValueError, match="inconsistent numbers of samples"):
        export_tariff(
            glm,
            tmp_path / "t.xlsx",
            X=df[feats],
            y=df["claim_amount"].iloc[:10],
            exposure_col="exposure",
        )


@pytest.mark.parametrize("n_columns", [1, 2])
@pytest.mark.parametrize("as_frame", [False, True])
def test_export_tariff_recalibration_rejects_2d_outcomes(
    tmp_path: Path, n_columns: int, as_frame: bool
) -> None:
    glm, df = _fit_glm(n=500)
    y = np.tile(df["claim_amount"].to_numpy()[:, None], (1, n_columns))
    if as_frame:
        y = pd.DataFrame(y)
    out = tmp_path / "invalid.xlsx"
    with pytest.raises(ValueError, match="one-dimensional"):
        export_tariff(glm, out, X=df[list(glm.feature_names_in_)], y=y, exposure_col="exposure")
    assert not out.exists()


# ---------------------------------------------------------------------------
# export_tariff (xlsx IO)
# ---------------------------------------------------------------------------


def test_export_tariff_writes_three_sheets(tmp_path: Path) -> None:
    glm, df = _fit_glm()
    out = tmp_path / "tariff.xlsx"
    p = export_tariff(
        glm,
        out,
        X=df[["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]],
        y=df["claim_amount"],
        exposure_col="exposure",
    )
    assert p == out and p.exists() and p.stat().st_size > 0
    xl = pd.read_excel(p, sheet_name=None)
    assert list(xl.keys()) == ["base_rate", "factors", "mappings"]


def test_export_tariff_recalibrated_reproduces_portfolio_total(tmp_path: Path) -> None:
    glm, df = _fit_glm()
    out = tmp_path / "tariff.xlsx"
    export_tariff(
        glm,
        out,
        X=df[["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]],
        y=df["claim_amount"],
        exposure_col="exposure",
        recalibrate=True,
    )
    xl = pd.read_excel(out, sheet_name=None)
    base = float(xl["base_rate"].iloc[0]["base_rate"])
    assert bool(xl["base_rate"].iloc[0]["recalibrated"]) is True

    # Rebuild the tariff dict from the sheets and verify it reproduces the total.
    factors = xl["factors"]
    numeric = {}
    categorical = {}
    for _, r in factors.iterrows():
        feat, lvl, f = r["feature"], r["level"], float(r["multiplicative_factor"])
        if lvl == "_per_unit":
            numeric[feat] = float(np.log(f))
        else:
            categorical.setdefault(feat, {})[lvl] = f
    tariff = {
        "base_rate": base,
        "reference": {
            f: next(lvl for lvl, fac in fs.items() if fac == 1.0) for f, fs in categorical.items()
        },
        "numeric": numeric,
        "categorical": categorical,
    }
    X = df[["driver_age", "vehicle_age", "region", "vehicle_brand"]]
    rate = apply_tariff(tariff, X)
    exp = df["exposure"].to_numpy(float)
    cam = df["claim_amount"].to_numpy(float)
    assert (rate * exp).sum() == pytest.approx(cam.sum(), rel=1e-6)


def test_export_tariff_no_recalibrate_stays_structural(tmp_path: Path) -> None:
    glm, df = _fit_glm()
    out = tmp_path / "tariff.xlsx"
    export_tariff(
        glm,
        out,
        X=df[["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]],
        y=df["claim_amount"],
        exposure_col="exposure",
        recalibrate=False,
    )
    xl = pd.read_excel(out, sheet_name=None)
    assert bool(xl["base_rate"].iloc[0]["recalibrated"]) is False
    # Structural base == extract_tariff base_rate.
    t = extract_tariff(glm)
    assert float(xl["base_rate"].iloc[0]["base_rate"]) == pytest.approx(t["base_rate"])


def test_export_tariff_recalibrate_requires_X_y_exposure(tmp_path: Path) -> None:
    glm, df = _fit_glm()
    out = tmp_path / "tariff.xlsx"
    with pytest.raises(ValueError, match="recalibrate=True requires"):
        export_tariff(glm, out, recalibrate=True)


def test_export_tariff_custom_reference_applied(tmp_path: Path) -> None:
    glm, df = _fit_glm()
    out = tmp_path / "tariff.xlsx"
    export_tariff(
        glm,
        out,
        X=df[["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]],
        y=df["claim_amount"],
        exposure_col="exposure",
        reference={"region": "urban", "vehicle_brand": "D"},
    )
    xl = pd.read_excel(out, sheet_name=None)
    m = xl["mappings"]
    assert str(m[m["feature"] == "region"].iloc[0]["reference_level"]) == "urban"
    assert str(m[m["feature"] == "vehicle_brand"].iloc[0]["reference_level"]) == "D"
    # And the chosen references have factor 1.0 in factors sheet.
    f = xl["factors"]
    urban = f[(f["feature"] == "region") & (f["level"] == "urban")].iloc[0]
    brandD = f[(f["feature"] == "vehicle_brand") & (f["level"] == "D")].iloc[0]
    assert float(urban["multiplicative_factor"]) == pytest.approx(1.0)
    assert float(brandD["multiplicative_factor"]) == pytest.approx(1.0)


# ---------------------------------------------------------------------------


def test_apply_tariff_rejects_nonfinite_numeric_input() -> None:
    glm, df = _fit_glm()
    tariff = extract_tariff(glm)
    X = df.head(3).copy()
    X.loc[X.index[0], "driver_age"] = np.nan
    with pytest.raises(ValueError, match="finite"):
        apply_tariff(tariff, X)


@pytest.mark.parametrize("strategy", ["quantile", "tree"])
def test_exported_intervals_independently_reproduce_boundary_premiums(
    tmp_path: Path, strategy
) -> None:
    from sklearn.pipeline import Pipeline

    from azoic.preprocessing import AutoBinner

    training = pd.DataFrame({"age": [0.0, 30.0, 60.0, 90.0] * 10, "exposure": 1.0})
    model = Pipeline(
        [
            ("binner", AutoBinner(cols=["age"], strategy=strategy, max_bins=3)),
            ("model", RiskGLM(family="poisson", link="log", alpha=0.01, exposure_col="exposure")),
        ]
    ).fit(training, np.tile([1.0, 10.0, 100.0, 100.0], 10))
    out = export_tariff(model, tmp_path / "boundaries.xlsx", recalibrate=False)
    sheets = pd.read_excel(out, sheet_name=None)
    factors = sheets["factors"]
    assert factors["feature"].unique().tolist() == ["age"]
    assert not sheets["base_rate"].loc[0, "recalibrated"]
    base_rate = float(sheets["base_rate"].loc[0, "base_rate"])
    bins = sheets["mappings"].query("role == 'binned'").iloc[0]
    edges = [float(value) for value in bins["mapping"].split(", ")]
    values = np.array(
        [-1e6]
        + [
            value
            for edge in edges
            for value in (np.nextafter(edge, -np.inf), edge, np.nextafter(edge, np.inf))
        ]
        + [1e6, np.nan]
    )
    rates = np.zeros(len(values))
    matches = np.zeros(len(values), dtype=int)
    assert bins["interval_closure"] == ("right" if strategy == "tree" else "left")
    interval_factors = sorted(
        [row for row in factors.itertuples() if row.level != "Missing"],
        key=lambda row: float(row.level[1:-1].split(", ")[0]),
    )
    bounds = [-np.inf, *edges, np.inf]
    for i, row in enumerate(interval_factors):
        assert row.level in bins["levels"]
        if bins["interval_closure"] == "right":
            mask = (values > bounds[i]) & (values <= bounds[i + 1])
        else:
            mask = (values >= bounds[i]) & (values < bounds[i + 1])
        if row.level == bins["missing_destination"]:
            mask |= np.isnan(values)
        matches += mask
        rates[mask] = base_rate * row.multiplicative_factor
    for row in factors.itertuples():
        if row.level == "Missing":
            mask = np.isnan(values)
            matches += mask
            rates[mask] = base_rate * row.multiplicative_factor

    np.testing.assert_array_equal(matches, np.ones(len(values), dtype=int))
    predictions = model.predict(pd.DataFrame({"age": values, "exposure": 1.0}))
    assert np.unique(np.round(predictions[np.isfinite(values)], 5)).size == 3
    np.testing.assert_allclose(rates, predictions, rtol=1e-10)


def test_export_tariff_unwraps_preprocessing_pipeline(tmp_path: Path) -> None:
    from azoic.workflow import ExperimentConfig, ModelSpec, run_experiment

    df = make_synthetic_portfolio(n=3000, seed=42)
    data = tmp_path / "portfolio.parquet"
    df.to_parquet(data)
    cfg = ExperimentConfig(
        data_path=str(data),
        spec={
            "target": "claim_amount",
            "exposure": "exposure",
            "claim_count": "claim_count",
        },
        features=["driver_age", "vehicle_age", "region", "vehicle_brand"],
        preprocessing={
            "binner": {"cols": ["driver_age", "vehicle_age"], "max_bins": 4},
            "grouper": {"cols": ["region", "vehicle_brand"], "strategy": "rare"},
        },
        models={
            "glm": ModelSpec(
                kind="glm",
                params={"family": "tweedie", "link": "log", "tweedie_power": 1.5},
            )
        },
    )
    _, estimators = run_experiment(cfg, return_estimators=True)
    pipeline = estimators["glm"]
    X = df[list(pipeline.feature_names_in_)]
    out = tmp_path / "pipeline-tariff.xlsx"

    export_tariff(
        pipeline,
        out,
        X=X,
        y=df["claim_amount"],
        exposure_col="exposure",
    )

    mappings = pd.read_excel(out, sheet_name="mappings")
    assert {"binned", "grouped"}.issubset(set(mappings["role"]))
    numeric = mappings.query("feature == 'driver_age' and role == 'binned'").iloc[0]
    assert numeric["interval_closure"] == "right"
    assert numeric["missing_destination"] in numeric["levels"]
    assert "Other" in mappings.loc[mappings["feature"] == "region", "levels"].iloc[0]

    transformed = X
    for _, step in pipeline.steps[:-1]:
        transformed = step.transform(transformed)
    tariff = extract_tariff(pipeline.named_steps["model"])
    tariff["base_rate"] = pd.read_excel(out, sheet_name="base_rate").loc[0, "base_rate"]
    predicted_total = float((apply_tariff(tariff, transformed) * df["exposure"]).sum())
    assert predicted_total == pytest.approx(df["claim_amount"].sum())


def test_distilled_pipeline_preserves_preprocessing_and_export_provenance(tmp_path: Path) -> None:
    from sklearn.pipeline import Pipeline

    from azoic.workflow import ExperimentConfig, ModelSpec, _split_indices, run_experiment

    df = make_synthetic_portfolio(n=3000, seed=42)
    data = tmp_path / "portfolio.parquet"
    df.to_parquet(data)
    cfg = ExperimentConfig(
        data_path=str(data),
        spec={
            "target": "claim_amount",
            "exposure": "exposure",
            "claim_count": "claim_count",
        },
        features=["driver_age", "vehicle_age", "region", "vehicle_brand"],
        preprocessing={
            "binner": {"cols": ["driver_age", "vehicle_age"], "max_bins": 4},
            "grouper": {"cols": ["region", "vehicle_brand"], "strategy": "rare"},
        },
        models={
            "gbm": ModelSpec(
                kind="gbm",
                params={
                    "objective": "tweedie",
                    "tweedie_variance_power": 1.5,
                    "n_estimators": 40,
                    "random_state": 42,
                },
            )
        },
    )
    _, estimators = run_experiment(cfg, return_estimators=True)
    teacher = estimators["gbm"]
    X = df[list(teacher.feature_names_in_)]
    train_idx, test_idx = _split_indices(cfg, df)

    student = distill_gbm(teacher, X.iloc[train_idx], X.iloc[test_idx])

    assert isinstance(student, Pipeline)
    assert student.steps[0][1] is not teacher.steps[0][1]
    out = tmp_path / "distilled.xlsx"
    export_tariff(student, out, X=X, recalibrate=False)
    sheets = pd.read_excel(out, sheet_name=None)
    base = sheets["base_rate"].iloc[0]
    assert list(sheets) == ["base_rate", "factors", "mappings"]
    assert base["distilled_from"] == "RiskGBM"
    assert base["teacher_objective"] == "tweedie"
    assert float(base["teacher_student_deviance"]) >= 0
    assert float(base["student_teacher_total_ratio"]) == pytest.approx(
        student.distillation_metrics_["student_teacher_total_ratio"]
    )
    assert {"binned", "grouped"}.issubset(set(sheets["mappings"]["role"]))

    transformed = student[:-1].transform(X.iloc[test_idx])
    tariff = extract_tariff(student.named_steps["model"])
    assert np.allclose(
        apply_tariff(tariff, transformed),
        student.predict(X.iloc[test_idx]),
    )


def test_m6_acceptance_tariff_reproduces_portfolio_total(tmp_path: Path) -> None:
    glm, df = _fit_glm(n=4000)
    out = tmp_path / "m6.xlsx"
    export_tariff(
        glm,
        out,
        X=df[["driver_age", "vehicle_age", "region", "vehicle_brand", "exposure"]],
        y=df["claim_amount"],
        exposure_col="exposure",
        recalibrate=True,
    )
    sheets = pd.read_excel(out, sheet_name=None)
    assert bool(sheets["base_rate"].iloc[0]["recalibrated"]) is True
    factors = sheets["factors"]
    assert set(factors["feature"]) == {
        "driver_age",
        "vehicle_age",
        "region",
        "vehicle_brand",
    }
    assert (factors["level"] == "_per_unit").sum() == 2
    assert len(factors[factors["level"] != "_per_unit"]) == 7
    mappings = sheets["mappings"]
    assert sorted(mappings["feature"]) == [
        "driver_age",
        "region",
        "vehicle_age",
        "vehicle_brand",
    ]


@pytest.mark.parametrize("with_frame", [False, True])
@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("kind", ["scaler", "function", "columns"])
def test_export_rejects_unsupported_preprocessing_before_writing(
    tmp_path, monkeypatch, kind, existing, with_frame
):
    from sklearn.compose import ColumnTransformer
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import FunctionTransformer, StandardScaler

    X = pd.DataFrame({"x": np.arange(12, dtype=float)})
    steps = {
        "scaler": StandardScaler().set_output(transform="pandas"),
        "function": FunctionTransformer(lambda frame: frame + 1),
        "columns": ColumnTransformer(
            [("x", "passthrough", ["x"])], verbose_feature_names_out=False
        ).set_output(transform="pandas"),
    }
    model = Pipeline(
        [("unsupported", steps[kind]), ("model", RiskGLM(family="poisson", link="log"))]
    ).fit(X, np.exp(X["x"] / 12))
    out = tmp_path / "tariff.xlsx"
    if existing:
        pd.DataFrame({"previous": [42]}).to_excel(out, index=False)
    before = out.read_bytes() if existing else None
    files = set(tmp_path.iterdir())

    def unexpected_writer(*args, **kwargs):
        raise AssertionError("unsupported preprocessing must fail before opening a workbook")

    monkeypatch.setattr(pd, "ExcelWriter", unexpected_writer)
    with pytest.raises(ValueError, match="unsupported preprocessing step 'unsupported'"):
        export_tariff(model, out, X=X if with_frame else None, recalibrate=False)
    assert set(tmp_path.iterdir()) == files
    assert (out.read_bytes() if out.exists() else None) == before


@pytest.mark.parametrize("recalibrate", [False, True])
def test_export_supported_pipeline_has_independent_workbook_prices(tmp_path, recalibrate):
    from ast import literal_eval

    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import FunctionTransformer

    from azoic.preprocessing import AutoBinner, AutoGrouper
    from azoic.workflow import _drop_columns

    X = pd.DataFrame(
        {
            "age": np.tile([20.0, 40.0, 60.0, 80.0, 90.0], 12),
            "segment": ["a", "a", "b", "c", "b"] * 12,
            "x": np.linspace(0, 1, 60),
            "exposure": 1.0,
            "claim_count": 1.0,
        }
    )
    y = np.exp(X["age"] / 100 + X["x"] / 10) * np.where(X["segment"] == "a", 2, 1)
    X["amount"] = y
    specials = {
        "exposure_col": "exposure",
        "claim_count_col": "claim_count",
        "target_col": "amount",
    }
    model = Pipeline(
        [
            ("skip", None),
            ("bins", AutoBinner(cols=["age"], strategy="quantile", max_bins=2, **specials)),
            ("pass", "passthrough"),
            ("groups", AutoGrouper(cols=["segment"], strategy="rare", min_exposure=20, **specials)),
            (
                "drop",
                FunctionTransformer(_drop_columns, kw_args={"columns": ["amount", "claim_count"]}),
            ),
            ("model", RiskGLM(family="poisson", link="log", exposure_col="exposure")),
        ]
    ).fit(X, y)
    original = X.copy(deep=True)
    out = tmp_path / "tariff.xlsx"
    pd.DataFrame({"previous": [42]}).to_excel(out, index=False)
    assert (
        export_tariff(
            model, out, X=X, y=X["amount"], exposure_col="exposure", recalibrate=recalibrate
        )
        == out
    )
    pd.testing.assert_frame_equal(X, original)
    assert list(tmp_path.iterdir()) == [out]
    sheets = pd.read_excel(out, sheet_name=None)
    assert list(sheets) == ["base_rate", "factors", "mappings"]
    factors, mappings = sheets["factors"], sheets["mappings"]
    assert set(factors["feature"]) == {"age", "segment", "x"}
    assert mappings.query("role == 'dropped'")["feature"].tolist() == ["amount", "claim_count"]
    bins = mappings.query("role == 'binned'").iloc[0]
    groups = mappings.query("role == 'grouped'").iloc[0]
    group_map = {
        literal_eval(source): literal_eval(destination)
        for item in groups["mapping"].split("; ")
        for source, destination in [item.split(" -> ")]
    }
    assert group_map["a"] == "a" and group_map["c"] == "Other"
    edge = float(bins["mapping"])
    probe = X.iloc[:5].copy()
    probe["age"] = [np.nextafter(edge, -np.inf), edge, np.nextafter(edge, np.inf), np.nan, 100.0]
    probe["segment"] = ["a", "c", "unseen", None, "b"]
    intervals = sorted(
        [level for level in factors.query("feature == 'age'")["level"] if level != "Missing"],
        key=lambda level: float(level[1:-1].split(", ")[0]),
    )
    assert bins["interval_closure"] == "left"
    age_levels = [
        intervals[0],
        intervals[1],
        intervals[1],
        bins["missing_destination"],
        intervals[1],
    ]
    prices = np.full(len(probe), float(sheets["base_rate"].loc[0, "base_rate"]))
    for row in factors.itertuples():
        if row.feature == "x":
            prices *= row.multiplicative_factor ** probe["x"].to_numpy()
        else:
            levels = (
                age_levels
                if row.feature == "age"
                else [
                    group_map.get(value, groups["unseen_destination"]) for value in probe["segment"]
                ]
            )
            prices[np.asarray(levels) == row.level] *= row.multiplicative_factor
    scale = X["amount"].sum() / np.dot(model.predict(X), X["exposure"]) if recalibrate else 1
    with pytest.warns(UserWarning, match="unseen levels"):
        expected = model.predict(probe) * scale
    np.testing.assert_allclose(prices, expected, rtol=1e-10)


@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("failure", ["serialization", "save", "replace"])
def test_export_failures_preserve_destination_and_remove_temporary_files(
    tmp_path, monkeypatch, existing, failure
):
    from openpyxl import Workbook

    model = RiskGLM(family="poisson", link="log").fit(
        pd.DataFrame({"x": np.arange(12, dtype=float)}), np.arange(1, 13, dtype=float)
    )
    out = tmp_path / "tariff.xlsx"
    if existing:
        pd.DataFrame({"previous": [42]}).to_excel(out, index=False)
    before = out.read_bytes() if existing else None
    files = set(tmp_path.iterdir())

    if failure == "serialization":
        to_excel = pd.DataFrame.to_excel

        def fail_serialization(self, *args, **kwargs):
            to_excel(self, *args, **kwargs)
            if kwargs["sheet_name"] == "factors":
                raise OSError("injected serialization failure")

        monkeypatch.setattr(pd.DataFrame, "to_excel", fail_serialization)
    elif failure == "save":

        def fail_save(self, filename):
            filename.write(b"incomplete workbook")
            raise OSError("injected save failure")

        monkeypatch.setattr(Workbook, "save", fail_save)
    else:

        def fail_replace(self, target):
            assert self.parent == out.parent and self != out
            assert pd.ExcelFile(self).sheet_names == ["base_rate", "factors", "mappings"]
            raise OSError("injected replace failure")

        monkeypatch.setattr(Path, "replace", fail_replace)
    with pytest.raises(OSError, match=f"injected {failure} failure"):
        export_tariff(model, out, recalibrate=False)
    assert set(tmp_path.iterdir()) == files
    assert (out.read_bytes() if out.exists() else None) == before


def test_distillation_preserves_generic_transforms_and_passthrough(tmp_path):
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import FunctionTransformer

    X = pd.DataFrame({"x": np.linspace(0, 10, 80), "exposure": 1.0})
    teacher = Pipeline(
        [
            ("skip", None),
            ("scale", FunctionTransformer(lambda frame: frame.assign(x=2 * frame["x"]))),
            ("pass", "passthrough"),
            ("model", RiskGBM(exposure_col="exposure", n_estimators=10, n_jobs=1)),
        ]
    ).fit(X.iloc[:60], np.exp(X["x"].iloc[:60] / 10))
    original = X.copy(deep=True)
    student = distill_gbm(teacher, X.iloc[:60], X.iloc[60:])
    pd.testing.assert_frame_equal(X, original)
    assert student.named_steps["scale"] is not teacher.named_steps["scale"]
    np.testing.assert_allclose(student[:-1].transform(X)["x"], 2 * X["x"])
    rates = student.predict(X.iloc[60:])
    assert rates.shape == (20,) and np.isfinite(rates).all() and (rates > 0).all()
    assert student.distillation_metrics_["student_teacher_total_ratio"] == pytest.approx(
        np.dot(rates, X["exposure"].iloc[60:])
        / np.dot(teacher.predict(X.iloc[60:]), X["exposure"].iloc[60:])
    )
    with pytest.raises(ValueError, match="unsupported preprocessing step 'scale'"):
        export_tariff(student, tmp_path / "unsupported.xlsx", recalibrate=False)
    assert not list(tmp_path.iterdir())
