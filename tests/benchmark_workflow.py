"""Opt-in Linux workflow measurement; generated evidence stays outside source control."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import subprocess
import sys
import threading
from contextlib import ExitStack, contextmanager
from datetime import UTC, datetime
from functools import wraps
from pathlib import Path
from time import perf_counter
from unittest.mock import patch

THREAD_ENV = ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS")


def rss_mib():
    return int(Path("/proc/self/statm").read_text().split()[1]) * os.sysconf("SC_PAGE_SIZE") / 2**20


class Meter:
    """Time disjoint stages and sample whole-process RSS, including native libraries."""

    def __init__(self):
        self.stages = {}
        self.active = None
        self.lock = threading.Lock()
        self.stop = threading.Event()

    def sample(self):
        with self.lock:
            if self.active is not None:
                row = self.stages[self.active]
                row["sampled_peak_rss_mib"] = max(row["sampled_peak_rss_mib"], rss_mib())

    def poll(self):
        while not self.stop.wait(0.01):
            self.sample()

    @contextmanager
    def running(self):
        thread = threading.Thread(target=self.poll, daemon=True)
        thread.start()
        try:
            yield self
        finally:
            self.stop.set()
            thread.join()

    def wrap(self, stage, function):
        @wraps(function)
        def measured(*args, **kwargs):
            with self.lock:
                if self.active is not None:
                    raise RuntimeError("benchmark stages must not overlap")
                row = self.stages.setdefault(
                    stage, {"calls": 0, "seconds": 0.0, "sampled_peak_rss_mib": 0.0}
                )
                self.active = stage
            self.sample()
            start = perf_counter()
            try:
                return function(*args, **kwargs)
            finally:
                row["seconds"] += perf_counter() - start
                row["calls"] += 1
                self.sample()
                with self.lock:
                    self.active = None

        return measured


def prepare(root, rows, threads):
    import numpy as np
    import pandas as pd
    import yaml

    from azoic import DatasetSpec, ExperimentConfig, ModelSpec, PreprocessingSpec

    rng = np.random.default_rng(42)
    numeric = [f"numeric_{i:02}" for i in range(20)]
    categorical = [f"category_{i:02}" for i in range(10)]
    data = {name: rng.uniform(-1, 1, rows) for name in numeric}
    data.update({name: pd.Categorical(rng.integers(0, 8, rows)) for name in categorical})
    frame = pd.DataFrame(data)
    frame["record_id"] = np.arange(rows)
    frame["period"] = np.arange(rows) * 24 // rows
    frame["exposure"] = rng.uniform(0.5, 1.0, rows)
    frequency = 0.15 * np.exp(0.4 * frame[numeric[0]] + 0.2 * frame[numeric[1]])
    frame["claim_count"] = rng.poisson(frequency * frame.exposure)
    frame["claim_amount"] = frame.claim_count * rng.gamma(3, 1000, rows)
    frame.to_parquet(root / "portfolio.parquet", index=False)
    config = ExperimentConfig(
        name="m62-synthetic-development",
        data_path="portfolio.parquet",
        spec=DatasetSpec(
            target="claim_amount",
            exposure="exposure",
            claim_count="claim_count",
            id_col="record_id",
            time_col="period",
            currency="EUR",
            claim_valuation_basis="simulated ultimate amounts",
            observation_period="synthetic periods 0–23",
            adjustments=["none; synthetic workload"],
        ),
        features=numeric + categorical,
        split="temporal",
        preprocessing=PreprocessingSpec(
            binner={"cols": numeric, "max_bins": 8},
            grouper={"cols": categorical, "max_groups": 4},
        ),
        models={
            "glm": ModelSpec(params={"family": "tweedie", "link": "log"}),
            "gbm": ModelSpec(
                kind="gbm",
                preprocess=False,
                params={
                    "objective": "tweedie",
                    "n_estimators": 50,
                    "num_leaves": 15,
                    "learning_rate": 0.05,
                    "n_jobs": threads,
                    "verbose": -1,
                },
            ),
        },
    )
    (root / "experiment.yaml").write_text(yaml.safe_dump(config.model_dump()), encoding="utf-8")
    return {"rows": rows, "numeric_predictors": 20, "categorical_predictors": 10, "seed": 42}


def measure(root, mode):
    import resource

    import numpy as np
    import pyarrow as pa
    from threadpoolctl import threadpool_info

    import azoic.workflow as workflow
    from azoic.models import RiskGBM, RiskGLM
    from azoic.preprocessing import AutoBinner, AutoGrouper
    from azoic.reporting import comparison_table, model_card
    from azoic.tariff import export_tariff

    threads = int(os.environ["OMP_NUM_THREADS"])
    pa.set_cpu_count(threads)
    pa.set_io_thread_count(threads)
    config = workflow.ExperimentConfig.from_yaml(root / "experiment.yaml")
    destination = root / mode
    destination.mkdir()
    meter = Meter()
    targets = [
        (workflow, "load_data", "loading"),
        (workflow, "_assemble_run", "assembly"),
        (workflow, "save_run", "persistence"),
        (workflow, "load_run", "persistence"),
        *[
            (workflow, name, "diagnostics")
            for name in ("gini", "op_ratio", "_tweedie_scores", "calibration_table")
        ],
        *[(cls, "fit", "preprocessing_fit") for cls in (AutoBinner, AutoGrouper)],
        *[(cls, "transform", "preprocessing_transform") for cls in (AutoBinner, AutoGrouper)],
        *[(cls, "fit", "model_fit") for cls in (RiskGLM, RiskGBM)],
        *[(cls, "predict", "prediction") for cls in (RiskGLM, RiskGBM)],
    ]
    start = perf_counter()
    with meter.running(), ExitStack() as stack:
        for owner, name, stage in targets:
            function = getattr(owner, name)
            if mode == "resume" and name == "fit":

                def reject(*args, **kwargs):
                    raise AssertionError("saved consumers must not fit")

                function = reject
            stack.enter_context(patch.object(owner, name, meter.wrap(stage, function)))
        frame = workflow.load_data(config.data_path, spec=config.spec)
        if mode == "fit":
            run, estimators = workflow.run_experiment(config, df=frame, return_estimators=True)
            workflow.save_run(run, estimators, root / "candidates.pkl")
        else:
            run, estimators = workflow.load_run(root / "candidates.pkl")
        card = meter.wrap("reporting", model_card)(run)
        (destination / "card.md").write_text(card, encoding="utf-8")
        meter.wrap("reporting", comparison_table)([run]).to_csv(
            destination / "comparison.csv", index=False
        )
        predictions = {}
        for name, estimator in estimators.items():
            rates = estimator.predict(frame[run.feature_names])
            if rates.shape != (len(frame),) or not np.all(np.isfinite(rates) & (rates > 0)):
                raise AssertionError("invalid benchmark rates")
            predictions[name] = rates
        np.savez(destination / "rates.npz", **predictions)
        meter.wrap("export", export_tariff)(
            estimators["glm"], destination / "tariff.xlsx", recalibrate=False
        )
    elapsed = perf_counter() - start
    return {
        "elapsed_seconds": elapsed,
        "stages": meter.stages,
        "unattributed_seconds": elapsed - sum(row["seconds"] for row in meter.stages.values()),
        "process_peak_rss_mib": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024,
        "threadpools": threadpool_info(),
        "arrow_cpu_threads": pa.cpu_count(),
        "arrow_io_threads": pa.io_thread_count(),
        "runtime_versions": run.runtime_versions,
        "data_fingerprint": run.data_fingerprint,
        "n_train": run.n_train,
        "n_validation": run.n_test,
    }


def verify(root):
    import numpy as np
    from openpyxl import load_workbook

    with np.load(root / "fit/rates.npz") as original, np.load(root / "resume/rates.npz") as saved:
        assert set(original.files) == set(saved.files) == {"glm", "gbm"}
        for name in original.files:
            np.testing.assert_array_equal(original[name], saved[name])
    for filename in ("card.md", "comparison.csv"):
        assert (root / "fit" / filename).read_bytes() == (root / "resume" / filename).read_bytes()
    books = [
        load_workbook(root / mode / "tariff.xlsx", read_only=True) for mode in ("fit", "resume")
    ]
    try:
        assert books[0].sheetnames == books[1].sheetnames == ["base_rate", "factors", "mappings"]
        for name in books[0].sheetnames:
            assert list(books[0][name].values) == list(books[1][name].values)
    finally:
        for book in books:
            book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True, help="new evidence directory")
    parser.add_argument("--rows", type=int, default=1_000_000)
    parser.add_argument("--threads", type=int, default=1)
    parser.add_argument(
        "--worker", choices=("prepare", "fit", "resume", "verify"), help=argparse.SUPPRESS
    )
    args = parser.parse_args()
    if sys.platform != "linux":
        parser.error("this RSS benchmark requires Linux /proc; no Windows performance claim")
    if args.rows < 500 or args.threads < 1:
        parser.error("rows must be at least 500 and threads at least 1")
    root = args.out.resolve()
    if args.worker:
        if args.worker == "prepare":
            result = prepare(root, args.rows, args.threads)
        elif args.worker == "verify":
            verify(root)
            return
        else:
            result = measure(root, args.worker)
        (root / f"{args.worker}.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        return
    root.mkdir(parents=True, exist_ok=False)
    environment = {
        **os.environ,
        **dict.fromkeys(THREAD_ENV, str(args.threads)),
        "MPLBACKEND": "Agg",
    }
    for phase in ("prepare", "fit", "resume", "verify"):
        subprocess.run(
            [
                sys.executable,
                str(Path(__file__).resolve()),
                "--out",
                str(root),
                "--rows",
                str(args.rows),
                "--threads",
                str(args.threads),
                "--worker",
                phase,
            ],
            env=environment,
            check=True,
        )
    fitted, saved = [json.loads((root / f"{mode}.json").read_text()) for mode in ("fit", "resume")]
    avoided = {
        stage: fitted["stages"][stage]["calls"] for stage in ("preprocessing_fit", "model_fit")
    }
    assert all(stage not in saved["stages"] for stage in avoided)
    metadata = {
        "recorded_at_utc": datetime.now(UTC).isoformat(),
        "platform": platform.platform(),
        "python": platform.python_version(),
        "cpu": next(
            line.split(":", 1)[1].strip()
            for line in Path("/proc/cpuinfo").read_text().splitlines()
            if line.startswith("model name")
        ),
        "affinity": sorted(os.sched_getaffinity(0)),
        "meminfo": Path("/proc/meminfo").read_text(),
        "thread_environment": {name: environment[name] for name in THREAD_ENV},
        "git_head": subprocess.check_output(["git", "rev-parse", "HEAD"], text=True).strip(),
        "git_status": subprocess.check_output(["git", "status", "--short"], text=True),
        "source_sha256": {
            str(path): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in (Path(__file__).resolve(), *Path("src/azoic").glob("*.py"))
        },
        "avoided_fit_calls": avoided,
        "avoided_fit_seconds": sum(fitted["stages"][stage]["seconds"] for stage in avoided),
        "consumer_outputs_equal": True,
    }
    (root / "environment.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "evidence": str(root),
                **{
                    key: metadata[key]
                    for key in (
                        "avoided_fit_calls",
                        "avoided_fit_seconds",
                        "consumer_outputs_equal",
                    )
                },
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
